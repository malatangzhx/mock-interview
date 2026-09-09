from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path

from docx import Document

QUESTION_KEYS = {"question", "题目", "问题", "prompt", "title", "name"}
FIELD_KEYS = {
    "topic": {"topic", "主题", "能力", "category"},
    "role": {"role", "岗位", "适用岗位"},
    "company": {"company", "公司"},
    "answer": {"answer", "答案", "参考答案", "answer_outline"},
    "source": {"source", "来源", "source_note"},
    "url": {"url", "链接", "source_url"},
    "interview_date": {"interview_date", "面试日期", "date"},
    "content_type": {"content_type", "内容类型"},
}


def _value(item, keys):
    for key, value in item.items():
        if str(key).strip().lower() in keys and value is not None:
            return "\n".join(map(str, value)) if isinstance(value, list) else str(value).strip()
    return ""


def parse_question_text(text: str) -> list[dict]:
    """Split explicit question boundaries; keep ambiguous multiline blocks for review."""
    entries, block, fence = [], [], False
    marker = re.compile(r"^(?:#{1,6}\s+(?=\S)|(?:\d+(?:\.\s+|[、)）])|[Qq]\d+[.:：]|问题\s*\d*[：:])\s*|[-•]\s+(?=\S))")
    def flush():
        if not any(x.strip() for x in block):
            block.clear()
            return
        raw = "\n".join(block).strip()
        lines = raw.splitlines()
        first = marker.sub("", lines[0], count=1)
        # Metadata is only recognized outside fenced code and after the question.
        fields, question, active, in_code = {}, [first], None, False
        for line in lines[1:]:
            if line.lstrip().startswith("```"):
                in_code = not in_code
            match = None if in_code else re.match(r"^(参考答案|答案|来源|公司|岗位|主题|链接|面试日期)\s*[：:]\s*(.*)$", line)
            if match:
                active = next(k for k, aliases in FIELD_KEYS.items() if match[1] in aliases)
                fields[active] = match[2]
            elif active:
                fields[active] += "\n" + line
            else:
                question.append(line)
        q = "\n".join(question).strip()
        entries.append({"question": q, **fields, "context": raw,
                        "issues": [] if re.search(r"[？?]|请|如何|什么|解释|说明", q) else ["无法确定是否为题目，请核对或删除"],
                        "position": f"条目 {len(entries)+1}"})
        block.clear()
    lines = text.replace("\r\n", "\n").splitlines()
    for line in lines:
        if line.lstrip().startswith("```"):
            fence = not fence
        if not fence and marker.match(line) and block:
            flush()
        # Unnumbered one-line questions are recognizable, continuations remain intact.
        elif not fence and block and re.match(r"^(请|如何|什么|为什么|怎样|解释|说明)", line) and block[-1].rstrip().endswith(("？", "?")):
            flush()
        block.append(line)
    flush()
    return entries


def parse_question_file(filename: str, data: bytes) -> list[dict]:
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("题库文件不能超过 5 MB")
    ext = Path(filename).suffix.lower()
    if ext not in {".txt", ".md", ".markdown", ".csv", ".json", ".docx"}:
        raise ValueError("题库支持 TXT、Markdown、CSV、JSON 和 DOCX")
    try:
        if ext == ".docx":
            if not zipfile.is_zipfile(io.BytesIO(data)):
                raise ValueError("DOCX 文件损坏或格式无效")
            doc = Document(io.BytesIO(data))
            entries = parse_question_text("\n".join(p.text for p in doc.paragraphs))
            for ti, table in enumerate(doc.tables):
                headers = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
                for ri, row in enumerate(table.rows[1:]):
                    item = dict(zip(headers, [c.text for c in row.cells]))
                    q = _value(item, QUESTION_KEYS)
                    entries.append({"question": q or "\n".join(c.text for c in row.cells),
                                    **{k: _value(item, keys) for k, keys in FIELD_KEYS.items()},
                                    "context": "\n".join(c.text for c in row.cells),
                                    "position": f"表格 {ti+1} 行 {ri+2}",
                                    "issues": [] if q else ["表格未识别题目列，请手工整理"]})
        else:
            text = data.decode("utf-8-sig")
            if ext in {".txt", ".md", ".markdown"}:
                entries = parse_question_text(text)
            else:
                payload = list(csv.DictReader(io.StringIO(text))) if ext == ".csv" else json.loads(text)
                if isinstance(payload, dict):
                    payload = next((payload[k] for k in ("questions", "items", "data") if k in payload), None)
                if not isinstance(payload, list):
                    raise ValueError("JSON 顶层应为数组，或包含 questions/items/data 数组")
                entries = []
                for i, item in enumerate(payload):
                    q = item if isinstance(item, str) else _value(item, QUESTION_KEYS) if isinstance(item, dict) else ""
                    entries.append({"question": q, **({k: _value(item, keys) for k, keys in FIELD_KEYS.items()} if isinstance(item, dict) else {}),
                                    "context": json.dumps(item, ensure_ascii=False), "position": f"数据行 {i+1}",
                                    "issues": [] if q else ["未识别题目字段，请手工整理"]})
    except (UnicodeDecodeError, json.JSONDecodeError, csv.Error) as exc:
        raise ValueError("文件格式无效；文本文件须使用 UTF-8 编码") from exc
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError("无法读取题库文件，请检查格式") from exc
    for entry in entries:
        entry["filename"] = Path(filename).name
        if not entry.get("content_type"):
            entry["content_type"] = "original"
    return entries
