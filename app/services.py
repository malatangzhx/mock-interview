from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .models import AnswerEvaluation, CandidateProfile, InterviewPlan, JobProfile, Requirement, Session, Topic, Turn


BUILTIN_JD = """AI应用开发实习生
负责使用 Python 开发大模型应用和后端接口；熟悉大模型 API、Prompt 设计、结构化输出与异常处理；
理解 RAG、Embedding、向量检索、召回与评估；了解 Agent、工具调用、数据存储、服务部署与基础软件工程；
能够清晰说明个人项目的目标、架构、贡献、技术取舍、难点和量化结果。"""

DATA_ANALYST_JD = """数据分析师
负责业务数据提取、清洗、分析和可视化，能够使用 SQL 与 Python 完成数据处理；
理解描述性统计、假设检验、指标体系和 A/B 测试，能够识别数据质量与口径问题；
熟悉 Excel 或主流 BI 工具，能够构建清晰的报表和仪表盘；
具备业务理解、问题拆解和沟通能力，能将分析结论转化为可执行建议并衡量结果。"""

BUILTIN_JOBS = {"ai": BUILTIN_JD, "data": DATA_ANALYST_JD}

SKILLS = ["Python", "FastAPI", "Flask", "Django", "Java", "C++", "RAG", "Agent", "Prompt", "大模型", "LLM", "向量数据库", "Milvus", "FAISS", "Chroma", "MySQL", "SQLite", "Redis", "Docker", "Git", "PyTorch", "LangChain"]

BASE_TOPICS = [
    {
        "key": "project",
        "name": "项目架构与个人贡献",
        "keywords": ["项目", "架构", "贡献", "难点", "结果"],
        "objective": "验证候选人能否清晰解释真实项目、个人边界和技术取舍",
        "question": "请选择一个最能代表你能力的项目，用两分钟说明目标、架构、你的具体贡献和最终结果。",
        "alternate": "如果由你重新做一次这个项目，你会保留和改变哪些架构决策？为什么？",
        "followups": ["请给出一个能证明你贡献的具体实现细节或结果数据。", "当时最重要的技术取舍是什么？还有哪些备选方案？"],
        "condition": "说明目标、架构、个人贡献、取舍和结果中的至少四项",
        "source": "岗位基线：项目经历结构化表达 Rubric",
    },
    {
        "key": "python",
        "name": "Python 与工程质量",
        "keywords": ["Python", "接口", "工程", "异常", "测试"],
        "objective": "验证 Python 工程实践、异常处理和可维护性意识",
        "question": "在一个调用外部大模型 API 的 Python 服务中，你会怎样设计超时、重试、限流和错误处理？",
        "alternate": "一个 Python AI 服务偶发超时且返回格式不稳定，你会如何定位并改造它？",
        "followups": ["重试可能带来哪些副作用，你会如何避免？", "你会记录哪些指标来判断系统是否健康？"],
        "condition": "覆盖超时、有限重试、错误分类和可观测性",
        "source": "Python/HTTP 服务可靠性通用工程实践",
    },
    {
        "key": "llm",
        "name": "大模型 API 与结构化输出",
        "keywords": ["LLM", "大模型", "API", "Prompt", "结构化"],
        "objective": "验证模型调用、Prompt 和输出校验的实践能力",
        "question": "如果模型必须稳定返回符合业务 Schema 的 JSON，你会如何设计提示、校验和失败恢复？",
        "alternate": "模型输出经常缺字段或类型错误时，你会如何在不掩盖问题的前提下提高可靠性？",
        "followups": ["仅在 Prompt 中要求 JSON 为什么还不够？", "如何做版本记录，使一次线上错误能够回放？"],
        "condition": "覆盖结构化约束、Schema 校验、有限重试和确定性降级",
        "source": "结构化输出与输入校验通用工程实践",
    },
    {
        "key": "rag",
        "name": "RAG 检索与评估",
        "keywords": ["RAG", "向量", "检索", "召回", "Embedding"],
        "objective": "验证对 RAG 全流程、召回质量和评估方法的理解",
        "question": "请设计一个面向课程资料问答的 RAG 流程，并说明你会如何评估检索和最终回答。",
        "alternate": "当 RAG 系统答案看似流畅却经常引用错资料时，你会怎样分层定位问题？",
        "followups": ["切分大小和召回数量分别会造成什么权衡？", "如果正确片段已召回但回答仍错误，你会检查什么？"],
        "condition": "覆盖切分、索引、召回、生成，以及检索与答案的分层评估",
        "source": "RAG 标准流程与检索评估通用 Rubric",
    },
    {
        "key": "agent",
        "name": "Agent 与工具调用",
        "keywords": ["Agent", "工具", "函数调用", "工作流"],
        "objective": "验证工具边界、状态管理和失败恢复设计",
        "question": "设计一个可以查课程信息并生成学习计划的 Agent，你会如何定义工具、状态和安全边界？",
        "alternate": "当 Agent 反复调用同一工具或产生危险参数时，你会如何约束执行？",
        "followups": ["哪些步骤更适合确定性工作流而不是交给模型决定？", "工具调用失败或返回脏数据时如何恢复？"],
        "condition": "覆盖工具 Schema、权限/校验、状态、循环限制和失败处理",
        "source": "工具调用安全与 Agent 工作流通用实践",
    },
    {
        "key": "backend",
        "name": "后端、数据与部署",
        "keywords": ["后端", "数据库", "部署", "Docker", "服务"],
        "objective": "验证接口、存储和服务部署的基础能力",
        "question": "将一个 AI 原型部署为可供同学试用的服务时，你会如何设计接口、数据存储、日志和部署流程？",
        "alternate": "一个本地 AI Demo 要变成多人可用服务，最先需要补齐哪些后端能力？",
        "followups": ["你会如何避免日志泄露简历和模型密钥？", "并发增加后最可能先出现什么瓶颈？"],
        "condition": "覆盖 API、存储、配置/密钥、日志和部署中的至少四项",
        "source": "Web 服务与隐私工程通用实践",
    },
    {
        "key": "sql",
        "name": "SQL 与数据提取",
        "keywords": ["SQL", "数据提取", "查询", "数据库", "窗口函数"],
        "objective": "验证 SQL 查询、聚合、连接与性能意识",
        "question": "如果要分析用户连续七天留存，你会如何定义口径并组织 SQL 查询？",
        "alternate": "一条包含多表连接和窗口函数的分析 SQL 很慢，你会怎样定位并优化？",
        "followups": ["如何处理重复记录和日期边界？", "请说明索引、执行计划和预聚合分别适合什么情况。"],
        "condition": "覆盖指标口径、数据关联、去重、时间窗口和校验",
        "source": "SQL 分析与指标口径通用 Rubric",
    },
    {
        "key": "statistics",
        "name": "统计分析与实验设计",
        "keywords": ["统计", "假设检验", "A/B", "实验", "显著性"],
        "objective": "验证统计推断、实验设计和结论边界意识",
        "question": "请设计一次新功能 A/B 测试，说明核心指标、分组、样本量和结果判断。",
        "alternate": "A/B 测试显示转化率显著提升，但收入下降，你会如何解释和继续分析？",
        "followups": ["统计显著是否等于业务上值得上线？", "如何避免多重检验和实验污染？"],
        "condition": "覆盖假设、指标、随机分组、样本量、显著性与业务意义",
        "source": "基础统计推断与在线实验通用 Rubric",
    },
    {
        "key": "quality",
        "name": "数据清洗与质量控制",
        "keywords": ["清洗", "数据质量", "异常值", "缺失值", "口径"],
        "objective": "验证对缺失、异常、重复和数据血缘问题的处理能力",
        "question": "拿到一份缺失、重复且指标口径不一致的业务数据，你会按什么顺序处理？",
        "alternate": "报表数据突然下降 30%，你会如何判断是真实业务变化还是数据链路问题？",
        "followups": ["缺失值为什么不能一律用均值填充？", "你会设置哪些自动化质量检查？"],
        "condition": "覆盖数据画像、异常定位、处理依据、口径确认和结果校验",
        "source": "分析数据质量管理通用实践",
    },
    {
        "key": "visualization",
        "name": "可视化与洞察表达",
        "keywords": ["可视化", "BI", "报表", "仪表盘", "Excel"],
        "objective": "验证图表选择、信息层级和面向决策者表达结论的能力",
        "question": "给管理层设计一页经营仪表盘时，你会展示什么，如何避免信息过载？",
        "alternate": "同一组数据面向管理层和运营人员汇报时，可视化应有哪些差异？",
        "followups": ["什么情况下不应该使用双轴图？", "你会如何让异常变化既醒目又不造成误导？"],
        "condition": "覆盖受众、核心指标、图表选择、对比基准和行动指引",
        "source": "数据可视化与信息设计通用 Rubric",
    },
    {
        "key": "business",
        "name": "业务分析与行动建议",
        "keywords": ["业务", "指标体系", "问题拆解", "建议", "转化"],
        "objective": "验证从业务问题到指标、分析和可执行建议的完整链路",
        "question": "某产品月活增长但付费率下降，你会如何拆解问题并形成行动建议？",
        "alternate": "业务方只说‘最近效果不好’，你会如何把它转化为可分析的问题？",
        "followups": ["如何区分相关关系与因果关系？", "分析建议上线后，你会怎样衡量它是否有效？"],
        "condition": "覆盖目标澄清、指标拆解、分群归因、验证与行动闭环",
        "source": "业务分析问题拆解通用 Rubric",
    },
]


def sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def extract_resume(filename: str, content_type: str | None, data: bytes) -> tuple[str, str]:
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("文件超过 5 MB 限制")
    ext = Path(filename).suffix.lower()
    if ext not in {".pdf", ".docx", ".txt"}:
        raise ValueError("仅支持 PDF、DOCX 和 TXT")
    try:
        if ext == ".pdf":
            if not data.startswith(b"%PDF"):
                raise ValueError("文件内容不是有效 PDF")
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("PDF 已加密，请粘贴简历文本")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            if not zipfile.is_zipfile(io.BytesIO(data)):
                raise ValueError("文件内容不是有效 DOCX")
            doc = Document(io.BytesIO(data))
            blocks = [p.text for p in doc.paragraphs]
            blocks.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            text = "\n".join(blocks)
        else:
            if b"\x00" in data[:1024]:
                raise ValueError("TXT 文件包含二进制内容")
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ValueError("TXT 不是 UTF-8 编码，请粘贴简历文本") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("文件损坏或无法解析，请粘贴简历文本") from exc
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 30:
        raise ValueError("未提取到足够的有效文本；扫描版 PDF 暂不支持，请粘贴文本")
    return text[:100_000], ext[1:]


def parse_candidate(text: str, document_id: str | None = None) -> CandidateProfile:
    lines = [x.strip(" •-\t") for x in text.splitlines() if x.strip()]
    skills = [skill for skill in SKILLS if re.search(re.escape(skill), text, re.I)]
    education = [line[:180] for line in lines if re.search(r"大学|学院|本科|硕士|研究生|博士|专业", line)][:4]
    projects = [line[:220] for line in lines if re.search(r"项目|系统|平台|助手|应用|研究", line)][:5]
    evidence = (education + projects + [line[:180] for line in lines if any(s.lower() in line.lower() for s in skills)])[:8]
    return CandidateProfile(
        education=education,
        skills=skills,
        projects=projects,
        evidence_refs=evidence,
        uncertain_fields=[] if skills and projects else ["项目或技能信息较少，请在确认页补充"],
        source_document_id=document_id,
        source_text_hash=sha(text),
    )


def parse_job(text: str, builtin: bool, target_company: str = "通用", use_user_question_bank: bool = False) -> JobProfile:
    raw = text or BUILTIN_JD
    title = next((x.strip() for x in raw.splitlines() if x.strip()), "自定义岗位")[:60]
    requirements = []
    for topic in BASE_TOPICS:
        hits = [k for k in topic["keywords"] if k.lower() in raw.lower()]
        if hits:
            requirements.append(Requirement(name=topic["name"], evidence="、".join(hits), priority=5))
    if not requirements:
        chunks = [x.strip() for x in re.split(r"[；。\n]", raw) if len(x.strip()) > 4][:6]
        requirements = [Requirement(name=x[:40], evidence=x[:120], priority=max(2, 5-i//2)) for i, x in enumerate(chunks)]
    return JobProfile(title=title, target_company=target_company.strip() or "通用",
                      use_user_question_bank=use_user_question_bank,
                      mode="validated" if builtin else "experimental", requirements=requirements[:8], source_text_hash=sha(raw))


def create_plan(candidate: CandidateProfile, job: JobProfile, question_bank: dict[str, list[dict]] | None = None,
                memberships: list[dict] | None = None, practiced_ids: set[str] | None = None,
                excluded_ids: set[str] | None = None) -> InterviewPlan:
    role = "data" if "数据分析" in job.title else "ai" if "ai" in job.title.lower() or "大模型" in job.title else "general"
    allowed = ({"project", "sql", "statistics", "quality", "visualization", "business", "python"} if role == "data"
               else {"project", "python", "llm", "rag", "agent", "backend"} if role == "ai" else {x["key"] for x in BASE_TOPICS})
    specs = {x["key"]: x for x in BASE_TOPICS}
    profile = " ".join(candidate.skills + candidate.projects).lower()
    req = " ".join(r.name + r.evidence for r in job.requirements).lower()
    practiced = practiced_ids or set()
    excluded = excluded_ids or set()
    notes, selected, seen, coverage = [], [], set(), Counter()
    bank_ids = list(dict.fromkeys(job.selected_bank_ids))
    memberships = memberships or []
    # Compatibility for old clients that selected the single historical bank.
    if job.use_user_question_bank and not bank_ids:
        bank_ids = list(dict.fromkeys(m["bank_id"] for m in memberships))
    all_items = [q for group in (question_bank or {}).values() for q in group]
    by_id = {q["id"]: q for q in all_items}

    def available(q):
        return q["id"] not in seen and q["id"] not in excluded and (job.allow_repeats or q["id"] not in practiced)

    def relevant(q):
        declared = q.get("user_role", "").lower()
        if declared and declared not in {"通用", "general", "全部"}:
            matches = (role == "data" and any(x in declared for x in ("data", "数据", "分析", "统计")) or
                       role == "ai" and any(x in declared for x in ("ai", "python", "大模型", "人工智能")) or
                       declared in job.title.lower())
            if not matches:
                return False
        if role == "data" and re.search(r"\bjava\b|jvm|spring|jmm", q["question"], re.I):
            return False
        return q.get("topic_key", "project") in allowed or q.get("topic_key") not in specs

    def relevance_score(q):
        spec = specs.get(q.get("topic_key"), specs["project"])
        return sum(2 * (k.lower() in req) + (k.lower() in profile) for k in spec["keywords"])

    def take(pool, count, kind):
        while len(selected) < count:
            candidates = [q for q in pool if available(q)]
            if not candidates:
                break
            q = min(candidates, key=lambda x: (x.get("bank_order", 0), coverage[x.get("topic_key", "project")],
                                              -relevance_score(x), x["id"] in practiced, x["id"]))
            seen.add(q["id"])
            # Content deduplication also covers matching baseline text.
            seen.add("text:" + sha(q["question"]))
            selected.append((q, kind))
            coverage[q.get("topic_key", "project")] += 1
            pool = [x for x in pool if "text:" + sha(x["question"]) not in seen]

    user_pool = []
    for m in memberships:
        if m["bank_id"] not in bank_ids or m["question_id"] not in by_id:
            continue
        q = {**by_id[m["question_id"]], "topic_key": m.get("topic") or by_id[m["question_id"]]["topic_key"],
             "user_role": m.get("role", ""), "membership": m, "bank_order": bank_ids.index(m["bank_id"])}
        if job.practice_mode == "normal" and not relevant(q):
            notes.append(f"普通模式未选用「{q['question'][:60]}」：与目标岗位不匹配；可在专项模式练习。")
            continue
        user_pool.append(q)
    take(user_pool, 6 if job.practice_mode == "specialized" else 4, "user")
    user_count = len(selected)
    if not job.allow_repeats and any(q["id"] in practiced for q in user_pool):
        notes.append("已展示过的用户题已跳过；开启重复练习后可重新选择。")
    if job.practice_mode == "normal":
        if bank_ids and user_count < 4:
            notes.append(f"相关可用用户题不足 4 道，本场选用 {user_count} 道，再以公司题和岗位通用题补充。")
        supplements = [q for q in all_items if q.get("source_claim") != "user_provided" and relevant(q) and not (role == "data" and q.get("topic_key") == "python")
                       and (role == "general" or role in q.get("roles", []) or "general" in q.get("roles", []))]
        company_pool = [q for q in supplements if job.target_company != "通用" and job.target_company in q.get("companies", [])]
        generic_pool = [q for q in supplements if not q.get("companies") or any(c in {"通用", "国内互联网通用"} for c in q.get("companies", []))]
        take(company_pool, 6, "supplement")
        take(generic_pool, 6, "supplement")
        baselines = [{"id": "baseline-"+sha(sp["question"])[:24], "question": sp["question"], "topic_key": sp["key"],
                      "alternate_question": sp["alternate"], "follow_up_questions": sp["followups"], "baseline": True}
                     for sp in BASE_TOPICS if sp["key"] in allowed and not (role == "data" and sp["key"] == "python") and "text:" + sha(sp["question"]) not in seen]
        take(baselines, 6, "supplement")
    if len(selected) < 6:
        notes.append(f"可用主问题不足 6 道，本场安排 {len(selected)} 道并缩短时长。")
    if not selected:
        notes.append("无可用新题，请选择其他题库、确认新题或主动开启重复练习。")
    topics = []
    for q, kind in selected:
        spec = specs.get(q.get("topic_key"), specs["project"])
        m = q.get("membership")
        is_company = job.target_company != "通用" and job.target_company in q.get("companies", [])
        note = (f"用户自行提供 · {m['bank_name']}" if m else
                f"{job.target_company}公开面经回忆题；{q.get('attribution', '')}" if is_company else
                q.get("attribution") or spec["source"])
        # Snapshots deliberately contain no reference answers before the interview.
        source_keys = ("original_question", "source", "url", "company", "role", "interview_date", "filename", "position", "modification_note", "confirmed_at")
        sources = ([{k: source.get(k, "") for k in source_keys} for source in m["sources"]] if m else
                   [{"source": note, "url": q.get("content_url", ""), "company": job.target_company if is_company else "",
                     "role": "", "interview_date": q.get("reported_at") or "", "filename": "", "position": q["id"]}])
        topics.append(Topic(name=spec["name"] if q.get("topic_key") in specs else q.get("topic_key", "用户题"), topic_key=q.get("topic_key", "project"),
            evidence=["用户题库选择" if m else "目标岗位匹配"], objective=spec["objective"] if not m else "说明该题的思路、依据和具体例子",
            weight=1 / max(1, len(selected)), question=q["question"], alternate_question=q.get("alternate_question", q["question"]),
            follow_up_questions=q.get("follow_up_questions") or ["请结合一个具体例子说明依据与结果。"],
            completion_condition=spec["condition"], estimated_minutes=6, source_level="C" if m or job.mode == "experimental" else q.get("authority_grade", "A"),
            source_note=note, question_bank_id=q["id"], question_id=q["id"], question_version=q.get("version", 1),
            source_url=(m.get("url") if m else q.get("content_url")) or None, source_license=q.get("license_name"),
            selection_kind=kind, selected_bank_id=m["bank_id"] if m else None, sources=sources,
            content_type=m.get("content_type", "original") if m else "recalled" if is_company else "generated_variant",
            selection_reason=f"题库顺序第 {bank_ids.index(m['bank_id'])+1} 位" if m else "公司匹配题" if is_company else "内置基线题" if q.get("baseline") else "岗位通用题"))
    return InterviewPlan(estimated_minutes=6*len(topics), topics=topics, strategy_version="question-bank-v2",
                         practice_mode=job.practice_mode, selected_bank_ids=bank_ids, allow_repeats=job.allow_repeats,
                         selection_notes=notes, user_question_count=user_count, supplement_question_count=len(topics)-user_count,
                         main_question_ids=[t.question_id for t in topics])


def evaluate(answer: str, topic: Topic) -> AnswerEvaluation:
    clean = re.sub(r"\s+", " ", answer).strip()
    length = len(clean)
    evidence_terms = ["例如", "项目", "我负责", "实现", "提升", "降低", "%", "用户", "数据", "结果"]
    depth_terms = ["因为", "原理", "权衡", "边界", "异常", "否则", "风险", "取舍", "监控", "评估"]
    struct_terms = ["首先", "其次", "最后", "第一", "第二", "一是", "二是"]
    topic_terms = [x for x in re.split(r"[、与和 ]", topic.name + " " + topic.objective) if len(x) > 1]
    relevance_hits = sum(1 for x in topic_terms if x.lower() in clean.lower())
    completeness = min(5, 1 + length // 90 + (1 if len(re.split(r"[，。；]", clean)) >= 5 else 0))
    depth = min(5, 1 + sum(1 for x in depth_terms if x in clean) // 2 + (1 if length > 180 else 0))
    evidence_score = min(5, 1 + sum(1 for x in evidence_terms if x in clean) // 2)
    relevance = min(5, 2 + relevance_hits)
    clarity = min(5, 2 + sum(1 for x in struct_terms if x in clean) + (1 if 60 <= length <= 700 else 0))
    correctness = None  # This heuristic does not verify any technical reference basis.
    confidence = round(min(0.9, 0.42 + min(length, 400) / 1000 + relevance_hits * 0.05), 2)
    strengths = []
    if relevance >= 4: strengths.append("回答紧扣当前问题")
    if evidence_score >= 3: strengths.append("提供了项目细节或结果证据")
    if depth >= 3: strengths.append("体现了原理或取舍思考")
    if not strengths: strengths.append("完成了基本作答")
    gaps = []
    if length < 80: gaps.append("回答偏短，可补充背景、行动和结果")
    if evidence_score < 3: gaps.append("缺少可核验的项目例子或量化结果")
    if depth < 3: gaps.append("可进一步说明原理、边界、取舍或异常场景")
    if relevance < 3: gaps.append("与问题核心的关联不够明确")
    return AnswerEvaluation(correctness=correctness, completeness=completeness, depth=depth, relevance=relevance, clarity=clarity, evidence=evidence_score, confidence=confidence, strengths=strengths, gaps=gaps or ["可进一步压缩表达并突出最关键结论"], evidence_quote=clean[:180])


def next_turn(session: Session, answer: str) -> tuple[Turn, str | None]:
    assert session.plan
    topic = session.plan.topics[session.current_topic_index]
    ev = evaluate(answer, topic)
    follow_count = session.topic_follow_ups.get(topic.id, 0)
    if ev.relevance <= 2 and follow_count < 2:
        action, reason, question = "clarify", "回答与问题核心关联较弱", f"请直接围绕“{topic.objective}”再说明一次，并给出一个具体例子。"
    elif ev.evidence <= 2 and follow_count < 2:
        action, reason, question = "request_evidence", "观点存在但缺少证据", topic.follow_up_questions[min(follow_count, len(topic.follow_up_questions)-1)]
    elif ev.depth <= 2 and follow_count < 2:
        action, reason, question = "probe_depth", "基本结论已有，但原理或边界不足", topic.follow_up_questions[min(follow_count, len(topic.follow_up_questions)-1)]
    else:
        action, reason, question = "switch_topic", "当前主题已完成有效验证或达到追问上限", None
    if action != "switch_topic":
        session.topic_follow_ups[topic.id] = follow_count + 1
    turn = Turn(question_id=topic.question_id, question_kind=(session.report or {}).get("pending_kind", "main"), sequence=len(session.turns)+1, topic_id=topic.id, topic_name=topic.name, question=session.report.pop("pending_question") if session.report and "pending_question" in session.report else topic.question, confirmed_answer=answer, evaluation=ev, next_action=action, decision_reason=reason)
    return turn, question


def make_report(session: Session) -> dict:
    assert session.plan and session.job
    groups: dict[str, list[Turn]] = {t.id: [] for t in session.plan.topics}
    for turn in session.turns:
        groups.setdefault(turn.topic_id, []).append(turn)
    topics = []
    weighted, weights = 0.0, 0.0
    for topic in session.plan.topics:
        turns = groups.get(topic.id, [])
        if not turns:
            topics.append({"name": topic.name, "status": "未考察", "score": None, "confidence": None, "source_level": topic.source_level, "evidence": "无", "strengths": [], "gaps": ["本场未覆盖，不代表能力不足"], "suggestion": "后续完成该主题的平行题"})
            continue
        evals = [x.evaluation for x in turns if x.evaluation]
        dimensions = ["completeness", "depth", "relevance", "clarity", "evidence"] + (["correctness"] if topic.answer_verification == "verified_used" else [])
        values = [getattr(e, d) for e in evals for d in dimensions if getattr(e, d) is not None]
        score = round(sum(values)/len(values), 1) if values else None
        confidence = round(sum(e.confidence for e in evals)/len(evals), 2) if evals else None
        if score is not None:
            weighted += score * topic.weight * (confidence or .5)
            weights += topic.weight * (confidence or .5)
        topics.append({"name": topic.name, "status": "已考察", "score": score, "confidence": confidence, "source_level": topic.source_level, "evidence": evals[-1].evidence_quote if evals else "回答已保存但未评分", "strengths": list(dict.fromkeys(x for e in evals for x in e.strengths))[:3], "gaps": list(dict.fromkeys(x for e in evals for x in e.gaps))[:3], "suggestion": "使用 STAR/结论—依据—边界结构重答，并补充一个可量化结果。"})
    for item, topic in zip(topics, session.plan.topics):
        item.update(question_id=topic.question_id, question=topic.question, sources=topic.sources,
                    selection_kind=topic.selection_kind, content_type=topic.content_type,
                    evaluation_basis="仅评价表达、结构和证据；未使用已核验答案，技术正确性不评分",
                    correctness=None, actual_questions=[{"question": t.question, "kind": t.question_kind} for t in groups.get(topic.id, [])])
    covered = sum(1 for x in topics if x["status"] == "已考察")
    for item in topics:
        baseline = session.baseline_scores.get(item["name"])
        item["baseline_score"] = baseline
        item["change"] = round(item["score"] - baseline, 1) if item["score"] is not None and baseline is not None else None
    return {"session_id": session.id, "job_title": session.job.title, "mode": session.job.mode, "phase": session.phase, "complete": not session.incomplete and covered == len(topics), "coverage": round(covered/len(topics), 2) if topics else 0, "overall_score": round(weighted/weights, 1) if weights else None, "score_note": "实验模式：技术正确性未经完整人工校准，不与正式岗位横向比较。" if session.job.mode == "experimental" else "五级训练量表，不代表真实招聘结论。低置信度评价已降权。", "topics": topics, "turn_count": len(session.turns), "model_version": session.model_version, "prompt_version": session.prompt_version}
